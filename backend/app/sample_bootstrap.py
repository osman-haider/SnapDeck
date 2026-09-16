"""Auto-generate the bundled sample templates on first run.

Rather than shipping opaque binary .pptx/.docx files in source control, we
build them programmatically (once, idempotently) so the project works
immediately after `pip install -r requirements.txt` with nothing to unzip
or download. Both templates use the same {{tag}} convention described in
templating.py and the backend README.
"""

from __future__ import annotations

from pathlib import Path

DECK_FILENAME = "template_deck.pptx"
ONEPAGER_FILENAME = "template_onepager.docx"


def ensure_sample_templates(sample_dir: Path) -> None:
    sample_dir.mkdir(parents=True, exist_ok=True)

    deck_path = sample_dir / DECK_FILENAME
    if not deck_path.exists():
        _build_sample_deck(deck_path)

    onepager_path = sample_dir / ONEPAGER_FILENAME
    if not onepager_path.exists():
        _build_sample_onepager(onepager_path)


def _build_sample_deck(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.63)

    # Slide 1 — Title
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = "{{title}}"
    title_slide.placeholders[1].text = "{{subtitle}}"

    # Slide 2 — Key metrics
    metrics_slide = prs.slides.add_slide(prs.slide_layouts[1])
    metrics_slide.shapes.title.text = "Key Metrics"
    body = metrics_slide.placeholders[1].text_frame
    body.text = "Revenue: {{revenue}}"
    for line in (
        "Growth: {{growth}}",
        "Top region: {{top_region}}",
        "Churn: {{churn}}",
        "NPS: {{nps}}",
    ):
        paragraph = body.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(20)

    # Slide 3 — Headline insight (AI-generated at render time)
    insight_slide = prs.slides.add_slide(prs.slide_layouts[1])
    insight_slide.shapes.title.text = "Headline Insight"
    insight_body = insight_slide.placeholders[1].text_frame
    insight_body.text = "{{ai_insight}}"
    insight_body.paragraphs[0].font.size = Pt(24)

    prs.save(str(path))


def _build_sample_onepager(path: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()

    doc.add_heading("{{title}}", level=1)
    subtitle = doc.add_paragraph("{{subtitle}}")
    subtitle.runs[0].italic = True

    doc.add_heading("Key Metrics", level=2)
    for line in (
        "Revenue: {{revenue}}",
        "Growth: {{growth}}",
        "Top region: {{top_region}}",
        "Churn: {{churn}}",
        "NPS: {{nps}}",
    ):
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Headline Insight", level=2)
    insight_paragraph = doc.add_paragraph("{{ai_insight}}")
    insight_paragraph.runs[0].font.size = Pt(13)

    doc.save(str(path))
