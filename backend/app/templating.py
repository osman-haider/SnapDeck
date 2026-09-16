"""Fill {{tag}} placeholders inside a .pptx or .docx template, and (optionally)
append a chart image as an extra slide/section.

Replacement happens at the *paragraph* level (joining every run's text,
doing the substitution on the joined string, then rewriting it into the
first run) rather than run-by-run. This matters because PowerPoint and Word
frequently split what looks like one `{{tag}}` across two or three separate
text runs behind the scenes (e.g. because of a spell-check squiggle or a
paste event) — replacing run-by-run would silently miss those. The tradeoff
is that a paragraph with mixed formatting collapses to the first run's
formatting after a substitution; for the deck/one-pager use case here
(mostly single-format lines) that's an acceptable, documented simplification.
"""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from docx import Document
from docx.shared import Inches as DocxInches
from pptx import Presentation
from pptx.util import Inches, Pt

TAG_PATTERN = re.compile(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}")


def _replace_in_paragraph(paragraph, fields: dict[str, str], used_tags: set[str]) -> None:
    full_text = "".join(run.text for run in paragraph.runs)
    if "{{" not in full_text:
        return

    def _substitute(match: re.Match) -> str:
        key = match.group(1)
        used_tags.add(key)
        return str(fields.get(key, match.group(0)))

    new_text = TAG_PATTERN.sub(_substitute, full_text)
    if new_text == full_text:
        return

    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = new_text


def _missing_fields(fields: dict[str, str], used_tags: set[str]) -> list[str]:
    return [key for key in fields if key not in used_tags and key != "ai_insight"]


# --------------------------------------------------------------------------- #
# PowerPoint (.pptx)
# --------------------------------------------------------------------------- #

def fill_pptx(template_bytes: bytes, fields: dict[str, str]) -> tuple[bytes, set[str], list[str]]:
    prs = Presentation(BytesIO(template_bytes))
    used_tags: set[str] = set()

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    _replace_in_paragraph(paragraph, fields, used_tags)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            _replace_in_paragraph(paragraph, fields, used_tags)

    missing = _missing_fields(fields, used_tags)
    out = BytesIO()
    prs.save(out)
    return out.getvalue(), used_tags, missing


def add_chart_slide_pptx(pptx_bytes: bytes, chart_png: bytes, title: str = "Chart") -> bytes:
    prs = Presentation(BytesIO(pptx_bytes))
    layout_index = 6 if len(prs.slide_layouts) > 6 else len(prs.slide_layouts) - 1
    slide = prs.slides.add_slide(prs.slide_layouts[layout_index])

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(8.8), Inches(0.7))
    title_frame = title_box.text_frame
    title_frame.text = title
    title_frame.paragraphs[0].font.size = Pt(22)
    title_frame.paragraphs[0].font.bold = True

    slide.shapes.add_picture(BytesIO(chart_png), Inches(1.0), Inches(1.3), width=Inches(8.0))

    out = BytesIO()
    prs.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------- #
# Word (.docx)
# --------------------------------------------------------------------------- #

def fill_docx(template_bytes: bytes, fields: dict[str, str]) -> tuple[bytes, set[str], list[str]]:
    doc = Document(BytesIO(template_bytes))
    used_tags: set[str] = set()

    for paragraph in doc.paragraphs:
        _replace_in_paragraph(paragraph, fields, used_tags)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_in_paragraph(paragraph, fields, used_tags)

    missing = _missing_fields(fields, used_tags)
    out = BytesIO()
    doc.save(out)
    return out.getvalue(), used_tags, missing


def add_chart_image_docx(docx_bytes: bytes, chart_png: bytes, title: str = "Chart") -> bytes:
    doc = Document(BytesIO(docx_bytes))
    doc.add_heading(title, level=2)
    doc.add_picture(BytesIO(chart_png), width=DocxInches(5.5))
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
